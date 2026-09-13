import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from mcp.server.fastmcp import FastMCP

# Инициализация MCP сервера с использованием FastMCP
mcp = FastMCP("File Exporter Server")

# Корневая папка для экспорта по умолчанию
BASE_EXPORT_DIR = Path.cwd() / "exported_docs"

def sanitize_filename(name: str) -> str:
    """Очищает и нормализует строку для безопасного использования в качестве имени файла."""
    cleaned = re.sub(r'[\\/*?:"<>|]', "", name)
    cleaned = re.sub(r'\s+', "_", cleaned).strip()
    return cleaned or "document"

@mcp.tool()
def export_to_file(
    title: str,
    content: str,
    format_type: str = "md",
    subfolder: str = ""
) -> str:
    try:
        target_dir = BASE_EXPORT_DIR
        if subfolder:
            safe_subfolder = sanitize_filename(subfolder)
            target_dir = target_dir / safe_subfolder
        
        target_dir.mkdir(parents=True, exist_ok=True)
        
        safe_title = sanitize_filename(title)
        ext = format_type.lower().strip(".")
        if ext not in ["md", "docx"]:
            ext = "md"
            
        file_path = target_dir / f"{safe_title}.{ext}"
        
        if ext == "md":
            file_content = f"# {title}\n\n{content}"
            file_path.write_text(file_content, encoding="utf-8")
            
        elif ext == "docx":
            doc = Document()
            
            # Настраиваем стиль "Normal" по умолчанию
            style_normal = doc.styles['Normal']
            style_normal.font.name = 'Times New Roman'
            style_normal.font.size = Pt(14)
            
            # Добавляем главный заголовок документа
            heading = doc.add_heading(title, level=1)
            heading.style.font.name = 'Times New Roman'
            
            lines = content.split("\n")
            table_rows = []
            
            def flush_table():
                """Вспомогательная функция для сборки накопившейся таблицы"""
                if not table_rows:
                    return
                # Фильтруем строки-разделители типа |---|---|
                filtered_rows = [r for r in table_rows if not all(c.strip("-") == "" for c in r)]
                if filtered_rows:
                    cols_count = len(filtered_rows[0])
                    table = doc.add_table(rows=len(filtered_rows), cols=cols_count)
                    table.style = 'Table Grid'
                    for r_idx, row_data in enumerate(filtered_rows):
                        for c_idx, cell_value in enumerate(row_data):
                            if c_idx < cols_count:
                                cell = table.cell(r_idx, c_idx)
                                cell.text = cell_value.strip()
                                # Форматируем текст в ячейках таблицы
                                for p in cell.paragraphs:
                                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                                    for run in p.runs:
                                        run.font.name = 'Times New Roman'
                                        run.font.size = Pt(12)
                table_rows.clear()
                doc.add_paragraph() # Пустой отступ после таблицы

            for line in lines:
                stripped = line.strip()
                
                # Если это строка таблицы Markdown
                if stripped.startswith("|") and stripped.endswith("|"):
                    cells = [c.strip() for c in stripped.strip("|").split("|")]
                    table_rows.append(cells)
                    continue
                else:
                    flush_table() # Если таблица закончилась, рендерим её
                
                # Обработка заголовков уровня 2 (##)
                if stripped.startswith("## "):
                    h = doc.add_heading(stripped[3:], level=2)
                    h.style.font.name = 'Times New Roman'
                # Обработка заголовков уровня 3 (###)
                elif stripped.startswith("### "):
                    h = doc.add_heading(stripped[4:], level=3)
                    h.style.font.name = 'Times New Roman'
                # Обработка маркированных списков (- или *)
                elif stripped.startswith("- ") or stripped.startswith("* "):
                    p = doc.add_paragraph(stripped[2:], style='List Bullet')
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                # Обычный текст / абзац
                elif stripped:
                    p = doc.add_paragraph(stripped)
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                else:
                    doc.add_paragraph()
            
            # На случай, если таблица была в самом конце файла
            flush_table()
            
            doc.save(str(file_path))
            
        abs_path = str(file_path.resolve())
        return f"Файл успешно создан.\nАбсолютный путь: {abs_path}"
        
    except Exception as e:
        return f"Произошла ошибка при создании файла: {str(e)}"

if __name__ == "__main__":
    mcp.run()