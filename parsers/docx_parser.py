from pathlib import Path
from docx import Document


def _cell(value: str) -> str:
    return str(value).replace("\r", " ").replace("\n", "<br>").replace("|", r"\|").strip()


def _table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    n = max(len(r) for r in rows)
    rows = [r + [""] * (n - len(r)) for r in rows]
    out = ["| " + " | ".join(_cell(x) for x in rows[0]) + " |",
           "| " + " | ".join("---" for _ in rows[0]) + " |"]
    out += ["| " + " | ".join(_cell(x) for x in row) + " |" for row in rows[1:]]
    return "\n".join(out)


def parse_docx(path: str | Path) -> str:
    doc = Document(path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for i, table in enumerate(doc.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md = _table(rows)
        if md:
            parts.append(f"Table {i}\n\n{md}")
    return "\n\n".join(parts).strip()
