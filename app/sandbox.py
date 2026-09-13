"""Session Sandbox Service for user uploaded files.

Handles secure isolated storage, multi-format parsing (PDF, DOCX, CSV, TXT, MD, Code),
and chat-scoped RAG retrieval for the Sber Meridian Assistant.
"""

import os
import re
import csv
import io
import json
import logging
from pathlib import Path
from typing import Dict, List, Optional, Any, Tuple

from agents.answer_agent import DocumentContext
from app.db import chat_history

logger = logging.getLogger(__name__)

SANDBOX_DIR = Path("data") / "sandbox_uploads"


def ensure_sandbox_dir(chat_id: str) -> Path:
    """Ensure sandbox directory exists for the given chat session."""
    chat_dir = SANDBOX_DIR / chat_id
    chat_dir.mkdir(parents=True, exist_ok=True)
    return chat_dir


def sanitize_filename(filename: str) -> str:
    """Sanitize filename to prevent path traversal."""
    clean = os.path.basename(filename)
    clean = re.sub(r'[^a-zA-Z0-9_\-\.\u0400-\u04FF]', '_', clean)
    return clean or "uploaded_file"


class SandboxService:
    """Manages file storage, parsing, and retrieval for chat-scoped sandbox."""

    def __init__(self):
        SANDBOX_DIR.mkdir(parents=True, exist_ok=True)

    def parse_file(self, file_bytes: bytes, filename: str) -> Tuple[str, str, Dict[str, Any]]:
        """
        Parse file content based on extension.
        Returns (text_content, file_type, parsed_meta).
        """
        ext = Path(filename).suffix.lower().lstrip(".")
        text_content = ""
        parsed_meta: Dict[str, Any] = {"extension": ext}

        if ext == "pdf":
            text_content, parsed_meta = self._parse_pdf(file_bytes)
        elif ext in ["docx", "doc"]:
            text_content, parsed_meta = self._parse_docx(file_bytes)
        elif ext in ["csv", "tsv"]:
            text_content, parsed_meta = self._parse_csv(file_bytes, delimiter="\t" if ext == "tsv" else None)
        else:
            # Code or text file (.txt, .md, .py, .json, .yaml, .yml, .sql, .js, .html, etc.)
            text_content, parsed_meta = self._parse_text(file_bytes, ext)

        return text_content.strip(), ext, parsed_meta

    def _parse_pdf(self, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """Extract text from PDF using PyMuPDF (fitz) or pypdf."""
        pages_text = []
        page_count = 0
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            page_count = len(doc)
            for page_num in range(page_count):
                page = doc[page_num]
                p_text = page.get_text("text") or ""
                if p_text.strip():
                    pages_text.append(f"--- Страница {page_num + 1} ---\n{p_text.strip()}")
            doc.close()
        except ImportError:
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                page_count = len(reader.pages)
                for page_num, page in enumerate(reader.pages):
                    p_text = page.extract_text() or ""
                    if p_text.strip():
                        pages_text.append(f"--- Страница {page_num + 1} ---\n{p_text.strip()}")
            except Exception as e:
                logger.warning(f"Failed to parse PDF with pypdf: {e}")
        except Exception as e:
            logger.warning(f"Failed to parse PDF with fitz: {e}")

        full_text = "\n\n".join(pages_text) if pages_text else "Текст из PDF не удалось извлечь."
        return full_text, {"page_count": page_count, "type": "document"}

    def _parse_docx(self, file_bytes: bytes) -> Tuple[str, Dict[str, Any]]:
        """Extract text and tables from DOCX using python-docx."""
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text.strip())

            for table in doc.tables:
                table_lines = []
                for row in table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    table_lines.append("| " + " | ".join(cells) + " |")
                if table_lines:
                    parts.append("\n" + "\n".join(table_lines) + "\n")

            full_text = "\n\n".join(parts)
            return full_text or "Документ DOCX пуст.", {"paragraph_count": len(doc.paragraphs), "type": "document"}
        except Exception as e:
            logger.warning(f"Failed to parse DOCX: {e}")
            return "Ошибка при чтении документа DOCX.", {"error": str(e), "type": "document"}

    def _parse_csv(self, file_bytes: bytes, delimiter: Optional[str] = None) -> Tuple[str, Dict[str, Any]]:
        """Parse CSV or TSV tabular data into text and table structure."""
        raw_text = self._decode_bytes(file_bytes)
        clean_text = raw_text.strip()
        if not clean_text:
            return "Пустой табличный файл.", {"type": "table", "headers": [], "rows": []}

        lines = [l for l in clean_text.splitlines() if l.strip()]
        if not delimiter:
            first_line = lines[0] if lines else ""
            if "\t" in first_line:
                delimiter = "\t"
            elif ";" in first_line:
                delimiter = ";"
            else:
                delimiter = ","

        try:
            reader = csv.reader(io.StringIO(clean_text), delimiter=delimiter)
            rows = list(reader)
            if rows:
                headers = [h.strip() for h in rows[0]]
                data_rows = [[c.strip() for c in r] for r in rows[1:]]

                # Format markdown table representation for LLM context
                md_lines = ["| " + " | ".join(headers) + " |"]
                md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for r in data_rows[:200]:  # Limit markdown representation to 200 rows for context window
                    padded = r + [""] * (len(headers) - len(r))
                    md_lines.append("| " + " | ".join(padded[:len(headers)]) + " |")

                md_text = "\n".join(md_lines)
                table_meta = {
                    "type": "table",
                    "headers": headers,
                    "rows": data_rows[:1000],
                    "total_rows": len(data_rows),
                }
                return md_text, table_meta
        except Exception as e:
            logger.warning(f"Failed to parse CSV: {e}")

        return clean_text, {"type": "table", "raw": True}

    def _parse_text(self, file_bytes: bytes, ext: str) -> Tuple[str, Dict[str, Any]]:
        """Decode plain text or code file."""
        text = self._decode_bytes(file_bytes)
        code_exts = ["py", "js", "ts", "json", "yaml", "yml", "sql", "html", "css", "xml", "asm", "sh", "bat"]
        meta_type = "code" if ext in code_exts else "text"
        return text, {"type": meta_type, "char_count": len(text)}

    @staticmethod
    def _decode_bytes(b: bytes) -> str:
        """Robustly decode byte string with encoding fallbacks."""
        for enc in ["utf-8", "utf-8-sig", "cp1251", "latin-1"]:
            try:
                return b.decode(enc)
            except UnicodeDecodeError:
                continue
        return b.decode("utf-8", errors="replace")

    def save_and_register_file(
        self,
        chat_id: str,
        filename: str,
        file_bytes: bytes,
    ) -> Dict[str, Any]:
        """Save file to disk and record in chat_sources database."""
        safe_name = sanitize_filename(filename)
        chat_dir = ensure_sandbox_dir(chat_id)

        import uuid
        source_id = str(uuid.uuid4())[:8]
        saved_filename = f"{source_id}_{safe_name}"
        file_path = chat_dir / saved_filename

        with open(file_path, "wb") as f:
            f.write(file_bytes)

        text_content, file_type, parsed_meta = self.parse_file(file_bytes, safe_name)

        record = chat_history.add_chat_source(
            chat_id=chat_id,
            filename=safe_name,
            file_path=str(file_path),
            file_type=file_type,
            size_bytes=len(file_bytes),
            text_content=text_content,
            parsed_meta=parsed_meta,
            source_id=source_id,
        )
        return record

    def search_sandbox_sources(
        self,
        chat_id: Optional[str],
        query: str,
        top_k: int = 3,
    ) -> List[DocumentContext]:
        """
        Search active files in the chat sandbox and return DocumentContext objects.
        """
        if not chat_id:
            return []

        sources = chat_history.list_chat_sources(chat_id, only_active=True)
        if not sources:
            return []

        query_lower = query.lower()
        query_words = set(re.findall(r'[a-zA-Zа-яА-Я0-9_]{3,}', query_lower))

        matched_docs: List[Tuple[float, Dict[str, Any]]] = []

        for src in sources:
            text = src.get("text_content", "")
            fname = src.get("filename", "")
            score = 0.5  # Base score for attached session document

            # If user mentions filename directly or partial name
            if fname.lower() in query_lower or any(p in query_lower for p in fname.lower().split(".") if len(p) > 2):
                score += 0.45

            # Word overlap with document text
            if query_words and text:
                text_lower = text[:5000].lower()
                matches = sum(1 for w in query_words if w in text_lower)
                overlap_ratio = matches / max(1, len(query_words))
                score += overlap_ratio * 0.35

            matched_docs.append((score, src))

        # Sort by relevance score
        matched_docs.sort(key=lambda x: x[0], reverse=True)

        results: List[DocumentContext] = []
        for score, src in matched_docs[:top_k]:
            sid = src["id"]
            fname = src["filename"]
            ftype = src["file_type"]
            content = src["text_content"]

            results.append(
                DocumentContext(
                    doc_id=f"sandbox-{sid}",
                    slug=f"sandbox-{sid}",
                    title=f"📎 {fname}",
                    product_name="Песочница",
                    product_code="ФАЙЛ",
                    section="Сессионный файл",
                    content=(
                        f"# Пользовательский файл: {fname}\n"
                        f"Тип файла: {ftype.upper()}\n\n"
                        f"{content}"
                    ),
                    attachment_path=src["file_path"],
                    attachment_format=ftype,
                    attachment_text=content if ftype in ["pdf", "docx", "csv", "tsv"] else None,
                    score=min(1.0, round(score, 2)),
                )
            )

        return results


sandbox_service = SandboxService()
